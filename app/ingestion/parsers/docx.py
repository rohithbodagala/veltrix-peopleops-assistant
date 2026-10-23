"""DOCX parsing with python-docx, walking the body in document order.

Paragraph styles are read from w:pStyle directly because `paragraph.style` is None for
some generated files. Tables are emitted where they appear. If a DOCX has almost no
text but embeds images (a scanned document saved as DOCX), the images are OCR'd.
"""
from __future__ import annotations

import re
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph

from app.ingestion.model import Block, Heading, Paragraph, ParsedDoc, Table
from app.ingestion.parsers.ocr import ocr_image_bytes, ocr_text_to_paragraphs


def _style_names(document) -> dict[str, str]:
    names: dict[str, str] = {}
    try:
        for style in document.styles:
            if style.style_id:
                names[style.style_id] = (style.name or style.style_id)
    except Exception:  # pragma: no cover - malformed styles part
        pass
    return names


def _paragraph_style(p_el, names: dict[str, str]) -> str:
    ppr = p_el.pPr
    if ppr is None or ppr.pStyle is None:
        return ""
    style_id = ppr.pStyle.val or ""
    return re.sub(r"\s+", "", names.get(style_id, style_id)).lower()


def _is_list(p_el) -> bool:
    ppr = p_el.pPr
    return ppr is not None and ppr.find(qn("w:numPr")) is not None


def _cell_text(cell) -> str:
    return re.sub(r"\s+", " ", " ".join(p.text for p in cell.paragraphs)).strip()


def _row_bold(row) -> bool:
    runs = [r for c in row.cells for p in c.paragraphs for r in p.runs if r.text.strip()]
    return bool(runs) and all(bool(r.bold) for r in runs)


def _table(tbl: DocxTable) -> Table | None:
    rows: list[list[str]] = []
    for row in tbl.rows:
        cells: list[str] = []
        seen: set[int] = set()
        for cell in row.cells:  # merged cells repeat; keep each once
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            cells.append(_cell_text(cell))
        if any(cells):
            rows.append(cells)
    if not rows:
        return None
    header = None
    if len(rows) > 1 and _row_bold(tbl.rows[0]):
        header = rows.pop(0)
    return Table(header=header, rows=rows)


def parse_docx(path: Path) -> ParsedDoc:
    document = docx.Document(str(path))
    names = _style_names(document)
    parsed = ParsedDoc(blocks=[])
    blocks: list[Block] = parsed.blocks
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            para = DocxParagraph(child, document)
            text = re.sub(r"\s+", " ", para.text).strip()
            if not text:
                continue
            style = _paragraph_style(child, names)
            m = re.match(r"heading(\d)", style)
            if style == "title":
                blocks.append(Heading(level=1, text=text))
            elif m:
                blocks.append(Heading(level=int(m.group(1)) + 1, text=text))
            elif "list" in style or _is_list(child):
                blocks.append(Paragraph(text="• " + text))
            else:
                blocks.append(Paragraph(text=text))
        elif child.tag == qn("w:tbl"):
            table = _table(DocxTable(child, document))
            if table is not None:
                blocks.append(table)

    text_len = sum(len(b.text) for b in blocks if isinstance(b, (Paragraph, Heading)))
    if text_len < 100:
        images = [
            rel.target_part.blob
            for rel in document.part.rels.values()
            if "image" in rel.reltype and not rel.is_external
        ]
        for blob in images:
            parsed.blocks.extend(ocr_text_to_paragraphs(ocr_image_bytes(blob), page=None))
            parsed.ocr_used = True
    return parsed
